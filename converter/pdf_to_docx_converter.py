from __future__ import annotations
import os
import io
import copy
import tempfile
import fitz  # PyMuPDF
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pdf2docx import Converter

from .utils import (
    get_logger,
    validate_input_file,
    ensure_output_path,
    ConversionError,
)
from .rtl_reconstruction import (
    page_arabic_script_ratio,
    reconstruct_page_paragraphs,
)

logger = get_logger(__name__)

_whitespace_patch_applied = False
_ocr_engine = None  # Global lazy OCR instance


def _get_ocr_engine():
    """Lazily initialize and cache the RapidOCR engine to avoid reloading models on every page."""
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            logger.error(f"Failed to initialize RapidOCR engine: {e}")
            _ocr_engine = False
    return _ocr_engine if _ocr_engine is not False else None


def _patch_pdf2docx_whitespace_bug():
    """Monkeypatch pdf2docx to prevent it from discarding whitespace-only spans."""
    global _whitespace_patch_applied
    if _whitespace_patch_applied:
        return

    try:
        import pdf2docx.text.Spans as spans_module
        from pdf2docx.text.TextSpan import TextSpan
        from pdf2docx.image.ImageSpan import ImageSpan

        def patched_restore(self, raws):
            for raw_span in raws:
                if "image" in raw_span:
                    span = ImageSpan(raw_span)
                else:
                    span = TextSpan(raw_span)
                    if not span.text and not span.style:
                        span = None
                self.append(span)
            return self

        spans_module.Spans.restore = patched_restore
        _whitespace_patch_applied = True
        logger.info("Applied pdf2docx whitespace-preservation patch.")
    except Exception as e:
        logger.warning(f"Could not patch pdf2docx whitespace: {e}")


def _append_docx_elements(master_doc: Document, sub_doc_path: str):
    """Merges sub-documents into master document, re-linking image assets and relationship IDs."""
    sub_doc = Document(sub_doc_path)
    rel_id_map = {}

    # Copy image parts from sub_doc into master_doc and create relationship ID mapping
    for rel_id, rel in sub_doc.part.rels.items():
        if "image" in rel.target_ref or "image" in rel.reltype:
            image_bytes = rel.target_part.blob
            image_stream = io.BytesIO(image_bytes)
            
            # get_or_add_image returns (r_id, image_object)
            r_id, _ = master_doc.part.get_or_add_image(image_stream)
            rel_id_map[rel_id] = r_id

    if len(master_doc.paragraphs) > 0 and master_doc.paragraphs[-1].text.strip() != "":
        master_doc.add_page_break()

    for element in sub_doc.element.body:
        if element.tag.endswith('sectPr'):
            continue

        new_elem = copy.deepcopy(element)

        # Remap image relationship attributes in XML (r:embed and r:id)
        for node in new_elem.xpath('.//*[@r:embed or @r:id]'):
            for attr in [
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',
                '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
            ]:
                if attr in node.attrib:
                    old_rid = node.attrib[attr]
                    if old_rid in rel_id_map:
                        node.attrib[attr] = rel_id_map[old_rid]

        master_doc.element.body.append(new_elem)


def _build_arabic_docx(pdf_page, output_docx_path):
    """Builds a temporary DOCX for a single Arabic page using existing RTL logic."""
    doc = Document()
    paragraphs = reconstruct_page_paragraphs(pdf_page)
    
    for para_data in paragraphs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement("w:bidi"))

        run = p.add_run(para_data["text"])
        run.font.size = Pt(para_data["size"])
        rPr = run._r.get_or_add_rPr()
        rPr.append(OxmlElement("w:rtl"))
        rFonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), para_data["font"])
        rPr.append(rFonts)
        
    doc.save(output_docx_path)


def _build_ocr_docx(pdf_page, output_docx_path):
    """Fallback engine for scanned pages using in-memory OCR & rendering."""
    doc = Document()
    engine = _get_ocr_engine()
    ocr_success = False

    if engine is not None:
        try:
            pix = pdf_page.get_pixmap(dpi=300)
            img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            img_np = np.array(img)
            
            result, _ = engine(img_np)
            
            if result:
                for item in result:
                    text = item[1]
                    if text.strip():
                        doc.add_paragraph(text)
                ocr_success = True
            else:
                doc.add_paragraph("[Scanned image page — no readable text detected]")
                ocr_success = True
                
        except Exception as e:
            logger.warning(f"OCR processing failed, falling back to image embedding: {e}")

    if not ocr_success:
        pix = pdf_page.get_pixmap(dpi=150)
        img_bytes = io.BytesIO(pix.tobytes("png"))
        doc.add_picture(img_bytes, width=Inches(6.0))

    doc.save(output_docx_path)


def convert_pdf_to_docx(
    input_path: str,
    output_path: str | None = None,
    start_page: int = 0,
    end_page: int | None = None,
    reconstruct_matrices_as_table: bool = False,
    progress_callback=None,
) -> str:
    """
    Enhanced hybrid routing pipeline:
    1. Arabic/RTL Pages -> rtl_reconstruction.py
    2. Scanned/Zero-Text Pages -> Pure In-Memory OCR Engine
    3. Standard Text Pages -> pdf2docx (monkeypatched)
    """
    validate_input_file(input_path, ".pdf")
    output_path = ensure_output_path(input_path, output_path, ".docx")

    _patch_pdf2docx_whitespace_bug()

    pdf = fitz.open(input_path)
    try:
        total_pages = len(pdf)
        last_page = end_page if end_page is not None else total_pages
        page_indices = list(range(max(0, start_page), min(last_page, total_pages)))
        
        if not page_indices:
            raise ConversionError("No pages in the requested range.")

        master_docx = Document()
        if master_docx.paragraphs:
            master_docx.paragraphs[0].clear()

        logger.info("Starting PDF->DOCX 3-way hybrid engine: %s -> %s", input_path, output_path)

        with tempfile.TemporaryDirectory() as temp_dir:
            for count, page_index in enumerate(page_indices):
                page = pdf[page_index]
                temp_pdf_path = os.path.join(temp_dir, f"page_{page_index}.pdf")
                temp_docx_path = os.path.join(temp_dir, f"page_{page_index}.docx")
                
                raw_text = page.get_text("text").strip()
                ratio = page_arabic_script_ratio(page)
                
                if ratio > 0.5:
                    logger.info(f"Page {page_index + 1}: Routing to RTL Geometry Engine")
                    _build_arabic_docx(page, temp_docx_path)
                elif len(raw_text) < 20:
                    logger.info(f"Page {page_index + 1}: Routing to In-Memory OCR Engine")
                    _build_ocr_docx(page, temp_docx_path)
                else:
                    logger.info(f"Page {page_index + 1}: Routing to pdf2docx Engine")
                    single_page_doc = fitz.open()
                    single_page_doc.insert_pdf(pdf, from_page=page_index, to_page=page_index)
                    single_page_doc.save(temp_pdf_path)
                    single_page_doc.close()
                    
                    cv = Converter(temp_pdf_path)
                    cv.convert(
                        temp_docx_path, 
                        start=0, 
                        end=1, 
                        parse_stream_table=reconstruct_matrices_as_table
                    )
                    cv.close()
                
                if os.path.exists(temp_docx_path):
                    _append_docx_elements(master_docx, temp_docx_path)
                    
                if progress_callback:
                    progress_callback(int(100 * (count + 1) / len(page_indices)))

        master_docx.save(output_path)
        logger.info("PDF->DOCX completed: %s", output_path)

    except Exception as e:
        logger.exception("PDF->DOCX conversion failed")
        raise ConversionError(f"PDF to Word conversion failed: {e}") from e
    finally:
        pdf.close()

    if not os.path.isfile(output_path) or os.path.getsize(output_path) == 0:
        raise ConversionError("Conversion reported success but no valid output file was produced.")

    return output_path